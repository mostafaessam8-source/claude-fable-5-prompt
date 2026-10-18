"""Build the new SAR sheets (standalone workbook) from work/model.json + the weekly workbook.
The standalone file also holds a value copy of 'Weekly Report Updates' so LibreOffice can recalc;
that copy is dropped when the sheets are moved into the weekly workbook (merge.py)."""
import sys, json, datetime, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL

model_p, wk_p, out_p = sys.argv[1:4]
M = json.load(open(model_p))

TEAL, DARK, GREY, LIGHT = "00778B", "3D3935", "A6A6A6", "E6F2F4"
F = "Arial"
def font(**k): k.setdefault("name", F); k.setdefault("size", 10); return Font(**k)
def fill(c): return PatternFill("solid", start_color=c, end_color=c)
thin = Side(style="thin", color="D9D9D9")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FT_IN, FT_F, FT_LINK = font(color="0000FF"), font(color="000000"), font(color="008000")
FILL_IN, FILL_KEY = fill("FFF9DB"), fill("FFFF00")
NUM, PCT, DATE, MON = '#,##0;(#,##0);"-"', '0.0%;(0.0%);"-"', 'dd-mmm-yyyy', 'mmm-yyyy'
MLN = '#,##0.0,,"M";(#,##0.0,,"M");"-"'
WR = "'Weekly Report Updates'"
WRA, WRH = WR + "!$A$1:$FZ$500", WR + "!$A$1:$FZ$1"
WCODE = "INDEX(%s,0,MATCH(\"Project Code\",%s,0))" % (WRA, WRH)
def wk_lookup(key_ref, hdr_ref):
    return '=IFERROR(INDEX(%s,MATCH(%s,%s,0),MATCH(%s,%s,0)),"")' % (WRA, key_ref, WCODE, hdr_ref, WRH)

def ddate(s): return datetime.datetime.strptime(s, "%Y-%m-%d") if s else None

wb = openpyxl.Workbook()
wb.remove(wb.active)
ORDER = ["Read Me", "Dashboard", "Projects", "Monthly Plan", "KPI Scorecard", "Delivery KPI", "Data Checks", "Review Notes", "Data Sources", "Settings"]
S = {n: wb.create_sheet(n) for n in ORDER}
for n in ORDER:
    S[n].sheet_properties.tabColor = TEAL if n in ("Projects", "Monthly Plan", "KPI Scorecard", "Delivery KPI", "Settings") else DARK
    S[n].sheet_view.showGridLines = False

def title(ws, text, sub, width_to=8):
    ws["A1"] = text; ws["A1"].font = font(size=16, bold=True, color=TEAL)
    ws["A2"] = sub; ws["A2"].font = font(size=9, italic=True, color="595959")
    ws.row_dimensions[1].height = 24

def link(cell, sheet, ref="A1", text=None):
    if text is not None: cell.value = text
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location="'%s'!%s" % (sheet, ref), display=str(cell.value))
    cell.font = font(color="0563C1", underline="single", size=cell.font.size if cell.font else 10)

def header_row(ws, row, cols, start_col=1):
    """cols: list of (header, kind, width, fmt, source_note)"""
    for i, (h, kind, w, fmt, note) in enumerate(cols):
        c = ws.cell(row, start_col + i, h)
        c.font = font(bold=True, color="FFFFFF"); c.border = BOX
        c.fill = fill(TEAL if kind == "in" else DARK if kind == "f" else "2E7D32")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[CL(start_col + i)].width = w
        n = ws.cell(row - 1, start_col + i, note)
        n.font = font(size=8, italic=True, color="0000FF" if kind == "in" else "008000" if kind == "wk" else "595959")
        n.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="center")
        n.fill = fill("F2F2F2")
    ws.row_dimensions[row].height = 42
    ws.row_dimensions[row - 1].height = 30

def put(ws, r, c, v, kind, fmt=None, wrap=False):
    cell = ws.cell(r, c, v)
    cell.font = FT_IN if kind == "in" else FT_LINK if kind == "wk" else FT_F
    if kind == "in": cell.fill = FILL_IN
    if fmt: cell.number_format = fmt
    cell.border = BOX
    if wrap: cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell

def add_table(ws, name, ref, style="TableStyleLight1"):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)

# ------------------------------------------------------------------ Settings
st = S["Settings"]
title(st, "Settings & Lists", "Yellow cells are the only settings to change. Lists feed the drop-downs in the input sheets.")
for c, w in zip("ABCDEFGHIJKLMNOP", (30, 22, 70, 3, 16, 3, 14, 3, 16, 3, 12, 3, 16, 46, 3, 14)):
    st.column_dimensions[c].width = w
for i, h in enumerate(("Setting", "Value", "What it does")):
    c = st.cell(3, 1 + i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(TEAL); c.border = BOX
SET = [
    ("Report date", "=MAX(%s!$B$2:$B$500)" % WR, DATE, "f", "Latest 'Report Date' in the weekly project files (Weekly Report Updates). Updates by itself on Refresh All."),
    ("Plan year", 2026, "0", "in", "Year of the Monthly Plan."),
    ("YTD cut-off month", datetime.datetime(2026, 9, 1), MON, "key", "Last month with confirmed actual spend. All YTD figures add the months up to and including this month. Change it once a month."),
    ("Next month", "=EDATE(B6,1)", MON, "f", "Month after the cut-off (used for the 'next month' plan columns)."),
    ("Weekly reports folder", r"D:\Final Dashboards\Weekly Reporty Updates\Weekly Report Updates", "@", "in",
     "Folder the Power Queries read (Data ▸ Refresh All). Info only: the path itself is inside the queries."),
    ("File status", "TEST VERSION – for discussion. Keep updating the old files until this file is approved.", "@", "in", "Shown on the Dashboard."),
]
for i, (lab, val, fmt, kind, note) in enumerate(SET):
    r = 4 + i
    st.cell(r, 1, lab).font = font(bold=True); st.cell(r, 1).border = BOX
    c = st.cell(r, 2, val); c.number_format = fmt; c.border = BOX
    c.font = FT_F if kind == "f" else FT_IN
    if kind == "key": c.fill = FILL_KEY
    elif kind == "in": c.fill = FILL_IN
    n = st.cell(r, 3, note); n.font = font(size=9, color="595959"); n.alignment = Alignment(wrap_text=True, vertical="top"); n.border = BOX
    st.row_dimensions[r].height = 30
for nm, ref in (("cfg_ReportDate", "$B$4"), ("cfg_Year", "$B$5"), ("cfg_YTDMonth", "$B$6"), ("cfg_NextMonth", "$B$7")):
    wb.defined_names[nm] = DefinedName(nm, attr_text="Settings!%s" % ref)

LISTS = [
    ("List_FundType", "Fund Type", "E", ["CAPEX", "PIF", "VRO", "PBU", "INFRA BU", "ISD BU", "Non-Profit"]),
    ("List_Program", "Program", "G", ["NSR", "ESR", "HHR"]),
    ("List_Phase", "Project Phase", "I", ["Planning", "Tendering", "Awarding", "Execution", "Closing", "On hold"]),
    ("List_Section", "Section", "K", ["FP", "CTP"]),
    ("List_Month", "Month", "M", [datetime.datetime(2026, m, 1) for m in range(1, 13)]),
]
LIST_REF = {}
for tname, h, col, vals in LISTS:
    c = st[col + "12"]; c.value = h; c.font = font(bold=True, color="FFFFFF"); c.fill = fill(TEAL)
    for i, v in enumerate(vals):
        x = st["%s%d" % (col, 13 + i)]; x.value = v; x.font = FT_IN; x.fill = FILL_IN; x.border = BOX
        if isinstance(v, datetime.datetime): x.number_format = MON
    add_table(st, tname, "%s12:%s%d" % (col, col, 12 + len(vals)))
    LIST_REF[tname] = "Settings!$%s$13:$%s$%d" % (col, col, 12 + len(vals) + 10)
st["A11"] = "Drop-down lists (add a value at the end of a list to make it available)"; st["A11"].font = font(bold=True, color=TEAL)
# glossary
st["N11"] = "Glossary"; st["N11"].font = font(bold=True, color=TEAL)
for i, h in enumerate(("Abbreviation", "Meaning")):
    c = st.cell(12, 14 + i, h); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(TEAL)
st.column_dimensions["N"].width = 16; st.column_dimensions["O"].width = 46
for i, (a, m) in enumerate(M["glossary"]):
    put(st, 13 + i, 14, a, "in"); put(st, 13 + i, 15, m, "in")
add_table(st, "Glossary", "N12:O%d" % (12 + len(M["glossary"])))
st.freeze_panes = "A4"

# ------------------------------------------------------------------ Projects
pj = S["Projects"]
title(pj, "Projects – one row per project (master data)",
      "Blue headers = type here · Dark headers = calculated, do not type · Green headers = read from the weekly project files (Refresh All). Row 3 says where each column comes from.")
MP = "'Monthly Plan'"
def mp(col): return "%s!$%s$5:$%s$2000" % (MP, col, col)
PCOLS = [  # key, header, kind, width, fmt, note
    ("code", "Code", "in", 9, "@", "Input · key"),
    ("short", "Short Name", "in", 16, None, "Input"),
    ("name", "Project Name", "in", 40, None, "Input"),
    ("program", "Program", "in", 9, None, "Input (list)"),
    ("section", "Section", "in", 8, None, "Input (list)"),
    ("fund", "Fund Type", "in", 11, None, "Input (list)"),
    ("po", "PO Number", "in", 10, "0", "Input"),
    ("kpi", "KPI Code", "in", 7, "0", "Input · KPI Scorecard code"),
    ("phase", "Project Phase", "in", 11, None, "Input (list)"),
    ("start", "Start Date", "in", 12, DATE, "Input"),
    ("end", "End Date", "in", 12, DATE, "Input"),
    ("full_cost", "Full Cost (SAR)", "in", 15, NUM, "Input"),
    ("cv", "Contract Value (SAR)", "in", 15, NUM, "Input"),
    ("wc_prev", "WC Previous Years (SAR)", "in", 15, NUM, "Input"),
    ("paid", "Paid from CV as per ERP (SAR)", "in", 15, NUM, "Input"),
    ("wcode", "Weekly Report Code", "in", 9, "@", "Input · code used in the weekly file"),
    ("remarks", "Remarks / Concern", "in", 40, None, "Input"),
    ("wc_cur", "WC Current Year (SAR)", "f", 14, NUM, "Calc: Monthly Plan actuals"),
    ("wc_tot", "Total WC (SAR)", "f", 14, NUM, "Calc: previous + current"),
    ("wc_pct", "Total WC %", "f", 8, PCT, "Calc: Total WC ÷ Contract"),
    ("wc_rem", "Remaining WC (SAR)", "f", 14, NUM, "Calc: Contract − Total WC"),
    ("pend", "Pending Payment (SAR)", "f", 14, NUM, "Calc: Total WC − Paid"),
    ("fy_b", "Budget Plan FY (SAR)", "f", 14, NUM, "Calc: Monthly Plan"),
    ("fy_v2", "V2 Forecast FY (SAR)", "f", 14, NUM, "Calc: Monthly Plan"),
    ("fy_f", "Contractor Forecast FY (SAR)", "f", 14, NUM, "Calc: Monthly Plan"),
    ("fy_var", "FY Variance Budget − V2 (SAR)", "f", 14, NUM, "Calc: Budget − V2"),
    ("ytd_b", "YTD Budget Plan (SAR)", "f", 14, NUM, "Calc: Monthly Plan ≤ cut-off"),
    ("ytd_v2", "YTD V2 Forecast (SAR)", "f", 14, NUM, "Calc: Monthly Plan ≤ cut-off"),
    ("ytd_a", "YTD Actual (SAR)", "f", 14, NUM, "Calc: Monthly Plan ≤ cut-off"),
    ("ytd_vb", "YTD Variance vs Budget (SAR)", "f", 14, NUM, "Calc: YTD Budget − Actual"),
    ("ytd_vv", "YTD Variance vs V2 (SAR)", "f", 14, NUM, "Calc: YTD V2 − Actual"),
    ("ytd_pct", "YTD % Variance vs Budget", "f", 9, PCT, "Calc: (Actual − Budget) ÷ Budget"),
    ("nm_v2", "Next Month V2 Forecast (SAR)", "f", 14, NUM, "Calc: Monthly Plan, next month"),
    ("nm_f", "Next Month Contractor Forecast (SAR)", "f", 14, NUM, "Calc: Monthly Plan, next month"),
    ("w_src", "Weekly Source File", "wk", 34, None, "Source.Name"),
    ("w_date", "Weekly Report Date", "wk", 12, DATE, "Report Date"),
    ("w_pl", "Planned % (cum)", "wk", 9, PCT, "Planned (%) - Cumulative"),
    ("w_ac", "Actual % (cum)", "wk", 9, PCT, "Actual (%) - Cumulative"),
    ("w_var", "Progress Variance", "f", 9, PCT, "Calc: Actual − Planned"),
    ("w_st", "Performance Status", "wk", 12, None, "Performance Status"),
    ("w_fin", "Forecast Finish", "wk", 12, DATE, "End Date\n(Forecast/Actual)"),
    ("w_spi", "SPI", "wk", 7, "0.00", "SPI"),
    ("chk", "Data Checks", "f", 22, None, "Calc: Data Checks sheet"),
]
PC = {k: CL(i + 1) for i, (k, *_r) in enumerate(PCOLS)}
header_row(pj, 4, [(h, kind, w, fmt, note) for (k, h, kind, w, fmt, note) in PCOLS])
P0 = 5
for i, p in enumerate(M["projects"]):
    r = P0 + i
    A = "$%s%d" % (PC["code"], r)
    f = {
        "wc_cur": "=SUMIFS(%s,%s,%s)" % (mp("G"), mp("A"), A),
        "wc_tot": "=%s%d+%s%d" % (PC["wc_prev"], r, PC["wc_cur"], r),
        "wc_pct": "=IFERROR(%s%d/%s%d,0)" % (PC["wc_tot"], r, PC["cv"], r),
        "wc_rem": "=%s%d-%s%d" % (PC["cv"], r, PC["wc_tot"], r),
        "pend": "=%s%d-%s%d" % (PC["wc_tot"], r, PC["paid"], r),
        "fy_b": "=SUMIFS(%s,%s,%s)" % (mp("D"), mp("A"), A),
        "fy_v2": "=SUMIFS(%s,%s,%s)" % (mp("F"), mp("A"), A),
        "fy_f": "=SUMIFS(%s,%s,%s)" % (mp("H"), mp("A"), A),
        "fy_var": "=%s%d-%s%d" % (PC["fy_b"], r, PC["fy_v2"], r),
        "ytd_b": '=SUMIFS(%s,%s,%s,%s,"<="&cfg_YTDMonth)' % (mp("D"), mp("A"), A, mp("B")),
        "ytd_v2": '=SUMIFS(%s,%s,%s,%s,"<="&cfg_YTDMonth)' % (mp("F"), mp("A"), A, mp("B")),
        "ytd_a": '=SUMIFS(%s,%s,%s,%s,"<="&cfg_YTDMonth)' % (mp("G"), mp("A"), A, mp("B")),
        "ytd_vb": "=%s%d-%s%d" % (PC["ytd_b"], r, PC["ytd_a"], r),
        "ytd_vv": "=%s%d-%s%d" % (PC["ytd_v2"], r, PC["ytd_a"], r),
        "ytd_pct": "=IFERROR((%s%d-%s%d)/%s%d,0)" % (PC["ytd_a"], r, PC["ytd_b"], r, PC["ytd_b"], r),
        "nm_v2": "=SUMIFS(%s,%s,%s,%s,cfg_NextMonth)" % (mp("F"), mp("A"), A, mp("B")),
        "nm_f": "=SUMIFS(%s,%s,%s,%s,cfg_NextMonth)" % (mp("H"), mp("A"), A, mp("B")),
        "w_var": '=IF(AND(ISNUMBER(%s%d),ISNUMBER(%s%d)),%s%d-%s%d,"")' % (PC["w_pl"], r, PC["w_ac"], r, PC["w_ac"], r, PC["w_pl"], r),
        "chk": "=IFERROR(INDEX('Data Checks'!$T$5:$T$200,MATCH(%s,'Data Checks'!$A$5:$A$200,0)),\"\")" % A,
    }
    for j, (k, h, kind, w, fmt, note) in enumerate(PCOLS):
        col = j + 1
        if kind == "in":
            v = p.get(k)
            if k in ("start", "end"): v = ddate(v)
            if k == "po" and v is not None: v = int(v)
            put(pj, r, col, v, "in", fmt)
        elif kind == "wk":
            put(pj, r, col, wk_lookup("$%s%d" % (PC["wcode"], r), "%s$3" % CL(col)), "wk", fmt)
        else:
            put(pj, r, col, f[k], "f", fmt)
PN = P0 + len(M["projects"]) - 1
add_table(pj, "Projects", "A4:%s%d" % (CL(len(PCOLS)), PN))
pj.freeze_panes = "C5"
for k, lst in (("program", "List_Program"), ("section", "List_Section"), ("fund", "List_FundType"), ("phase", "List_Phase")):
    dv = DataValidation(type="list", formula1="=" + LIST_REF[lst], allow_blank=True)
    dv.add("%s5:%s500" % (PC[k], PC[k])); pj.add_data_validation(dv)
dv = DataValidation(type="list", formula1="='KPI Scorecard'!$A$5:$A$40", allow_blank=True); dv.add("%s5:%s500" % (PC["kpi"], PC["kpi"])); pj.add_data_validation(dv)

# ------------------------------------------------------------------ Monthly Plan
mpw = S["Monthly Plan"]
title(mpw, "Monthly Plan 2026 – one row per project per month",
      "Type the monthly amounts in the blue columns. A cell showing =G<row> takes the Actual of that month (used for closed months). Cumulative and YTD columns are calculated.")
MCOLS = [
    ("Code", "in", 9, "@", "Input · Projects code"),
    ("Month", "in", 11, MON, "Input · 1st of month"),
    ("Project", "f", 16, None, "Calc: Projects"),
    ("Budget Plan (Incr)", "in", 14, NUM, "Input · Spend Plan as per Budgeting"),
    ("V0 Forecast Jan-26 (Incr)", "in", 14, NUM, "Input · shared with PC"),
    ("V2 Forecast Mar-26 (Incr)", "in", 14, NUM, "Input · shared with PC"),
    ("Actual Spend (Incr)", "in", 14, NUM, "Input · work confirmed"),
    ("Contractor Forecast (Incr)", "in", 14, NUM, "Input · cash flow / progress"),
    ("Invoice-related Activities", "in", 30, None, "Input"),
    ("Milestones", "in", 24, None, "Input"),
    ("Remarks / Concern", "in", 30, None, "Input"),
    ("In YTD?", "f", 7, None, "Calc: Month ≤ cut-off"),
    ("Budget Plan (Cum)", "f", 14, NUM, "Calc: running total"),
    ("V0 Forecast (Cum)", "f", 14, NUM, "Calc: running total"),
    ("V2 Forecast (Cum)", "f", 14, NUM, "Calc: running total"),
    ("Actual Spend (Cum)", "f", 14, NUM, "Calc: running total"),
    ("Contractor Forecast (Cum)", "f", 14, NUM, "Calc: running total"),
]
header_row(mpw, 4, MCOLS)
def mval(v, r):
    if v == "=ACT": return "=G%d" % r
    return v
for i, m in enumerate(M["monthly"]):
    r = 5 + i
    put(mpw, r, 1, m["code"], "in", "@"); put(mpw, r, 2, ddate(m["month"]), "in", MON)
    put(mpw, r, 3, '=IFERROR(INDEX(Projects!$B$5:$B$500,MATCH($A%d,Projects!$A$5:$A$500,0)),"Not in Projects")' % r, "f")
    for c, k in zip(range(4, 9), ("budget", "v0", "v2", "actual", "fcst")):
        put(mpw, r, c, mval(m[k], r), "in", NUM)
    for c, k in zip(range(9, 12), ("invoice", "milestone", "remarks")):
        put(mpw, r, c, m[k], "in", wrap=False)
    put(mpw, r, 12, '=IF($B%d<=cfg_YTDMonth,"Yes","No")' % r, "f")
    for c, src in zip(range(13, 18), "DEFGH"):
        put(mpw, r, c, '=SUMIFS($%s$5:$%s$2000,$A$5:$A$2000,$A%d,$B$5:$B$2000,"<="&$B%d)' % (src, src, r, r), "f", NUM)
MN = 4 + len(M["monthly"])
add_table(mpw, "Monthly_Plan", "A4:Q%d" % MN)
mpw.freeze_panes = "D5"
dv = DataValidation(type="list", formula1="=Projects!$A$5:$A$500", allow_blank=False); dv.add("A5:A3000"); mpw.add_data_validation(dv)
dv = DataValidation(type="list", formula1="=" + LIST_REF["List_Month"], allow_blank=False); dv.add("B5:B3000"); mpw.add_data_validation(dv)

# ------------------------------------------------------------------ KPI Scorecard
ks = S["KPI Scorecard"]
title(ks, "KPI Scorecard 2026",
      "Measured From: Monthly Plan = spend KPIs from the projects with that KPI Code · Weekly Report = calculated from the weekly project files (block under the table) · Manual = type target and actual in the blue cells.")
NF2 = "#,##0.00;(#,##0.00);\"-\""
KCOLS = [
    ("KPI Code", "in", 7, "0", "Input · key"),
    ("Perspective", "in", 18, None, "Input"),
    ("Objective / KPI", "in", 42, None, "Input"),
    ("KPI Weight", "in", 8, PCT, "Input"),
    ("Measured From", "in", 14, None, "Monthly Plan / Weekly Report / Manual"),
    ("FY Target / Plan", "in", 15, NF2, "Spend: Calc · else Input"),
    ("FY V2 Forecast", "in", 15, NF2, "Spend: Calc · else Input"),
    ("FY Variance", "f", 14, NF2, "Calc: V2 − Target"),
    ("YTD Target / Plan", "in", 15, NF2, "Spend: Calc · else Input"),
    ("YTD V2 Forecast", "in", 15, NF2, "Spend: Calc · else Input"),
    ("YTD Actual", "in", 15, NF2, "Monthly Plan / Weekly Report: Calc · Manual: Input"),
    ("YTD Variance", "f", 14, NF2, "Calc: Actual − Target"),
    ("% Achieved", "f", 9, PCT, "Calc: vs Target (max 100%)"),
    ("KPI Result", "f", 9, '0.00%;(0.00%);"-"', "Calc: % Achieved × Weight"),
    ("Indicator", "f", 11, None, "Calc: from % Achieved"),
    ("Remarks", "in", 30, None, "Input"),
    ("Linked Project", "in", 9, "@", "Input (Five Bridges KPI)"),
    ("Better When", "in", 9, None, "Input: Higher / Lower"),
    ("How the YTD Actual is found", "f", 52, None, "Source of the actual"),
]
header_row(ks, 4, KCOLS)
def pjr(k): return "Projects!$%s$5:$%s$500" % (PC[k], PC[k])
KN = 4 + len(M["kpis"])
# weekly-report block under the table (every number shows where it comes from)
B0 = KN + 5
def hdrq(h): return '"' + h.replace("\n", '"&CHAR(10)&"') + '"'
def wsum(*hs): return "+".join("SUM(INDEX(%s,0,MATCH(%s,%s,0)))" % (WRA, hdrq(h), WRH) for h in hs)
WBLK = [
    ("fat", "Fatalities (cumulative, all weekly files)", "=" + wsum("Fatality (Cumulative)"), "0", "Weekly Report Updates › Fatality (Cumulative)"),
    ("maj", "Major injuries = Lost Time + Restricted Work-Day cases", "=" + wsum("Lost Time Incident (Cumulative)", "Restricted Work-Day Case (Cumulative)"), "0", "Weekly Report Updates › Lost Time Incident (Cumulative) + Restricted Work-Day Case (Cumulative)"),
    ("min", "Minor injuries = Medical Treatment + First Aid cases", "=" + wsum("Medical Treatment Case (Cumulative)", "First Aid Case (Cumulative)"), "0", "Weekly Report Updates › Medical Treatment Case (Cumulative) + First Aid Case (Cumulative)"),
    ("nm", "Near misses", "=" + wsum("Near Misses (Cumulative)"), "0", "Weekly Report Updates › Near Misses (Cumulative)"),
    ("mh", "Man-hours worked", "=" + wsum("Total Man-Hours\n(Cumulative)"), "#,##0", "Weekly Report Updates › Total Man-Hours (Cumulative)"),
    ("fwi", "FWI = (fatalities + major/10 + minor/200 + near misses/1000) ÷ man-hours × 1,000,000", None, "0.000", "Formula from the 2026 Balanced Scorecard (Projects Department)"),
    ("ev", "Cumulative Earned Value (SAR)", "=" + wsum("Cumulative EV (SAR)"), NUM, "Weekly Report Updates › Cumulative EV (SAR)"),
    ("pv", "Cumulative Planned Value (SAR)", "=" + wsum("Cumulative PV (SAR)"), NUM, "Weekly Report Updates › Cumulative PV (SAR)"),
    ("spi", "SPI = Earned Value ÷ Planned Value (all weekly projects)", None, "0.00", "Formula from the 2026 Balanced Scorecard"),
    ("due", "Delivery projects due by the report date", "=COUNTIF('Delivery KPI'!$L$5:$L$200,\"Yes\")", "0", "Delivery KPI › Due by Report Date?"),
    ("dlv", "… of them delivered (weekly actual progress 100%)", "=COUNTIFS('Delivery KPI'!$L$5:$L$200,\"Yes\",'Delivery KPI'!$M$5:$M$200,\"Yes\")", "0", "Delivery KPI › Delivered?"),
    ("pdl", "% delivery = delivered ÷ due", None, PCT, "Formula from the 2026 Balanced Scorecard: total projects delivered / planned"),
]
WR_ = {k: "$F$%d" % (B0 + 1 + i) for i, (k, *_x) in enumerate(WBLK)}
c = ks.cell(B0, 3, "Weekly report inputs (calculated from the weekly project files – do not type)"); c.font = font(bold=True, size=11, color="2E7D32")
for i, (k, lab, f, fmt, src) in enumerate(WBLK):
    r = B0 + 1 + i
    if k == "fwi": f = "=IFERROR((%s+%s/10+%s/200+%s/1000)/%s*1000000,0)" % (WR_["fat"], WR_["maj"], WR_["min"], WR_["nm"], WR_["mh"])
    if k == "spi": f = "=IFERROR(%s/%s,0)" % (WR_["ev"], WR_["pv"])
    if k == "pdl": f = "=IFERROR(%s/%s,0)" % (WR_["dlv"], WR_["due"])
    a = ks.cell(r, 3, lab); a.font = font(size=9, bold=k in ("fwi", "spi", "pdl")); a.border = BOX; a.alignment = Alignment(wrap_text=True)
    put(ks, r, 6, f, "wk" if k not in ("fwi", "spi", "pdl") else "f", fmt)
    n = ks.cell(r, 9, src); n.font = font(size=8, italic=True, color="595959")
    ks.row_dimensions[r].height = 24
AUTO = [  # (name pattern, YTD actual formula, how, better, linked)
    (r"fatality weight", "=" + WR_["fwi"], "Weekly report: FWI from the safety figures (block below). Lower is better.", "Lower", None),
    (r"five bridges", '=IFERROR(INDEX(%s,MATCH($Q{r},%s,0)),0)' % (pjr("w_ac"), pjr("code")), "Weekly report: Actual % (cum) of the linked project (Projects ◂ weekly).", "Higher", "N003"),
    (r"schedule performance", "=" + WR_["spi"], "Weekly report: Σ Earned Value ÷ Σ Planned Value (block below).", "Higher", None),
    (r"delivery against approved business plan", "=" + WR_["pdl"], "Delivery KPI sheet: projects delivered ÷ projects due by the report date.", "Higher", None),
]
for i, k in enumerate(M["kpis"]):
    r = 5 + i
    auto = next((a for a in AUTO if re.search(a[0], k["name"], re.I)), None)
    better = auto[3] if auto else ("Lower" if re.search(r"incident weight", k["name"], re.I) else "Higher")
    mf = "Monthly Plan" if k["spend"] else "Weekly Report" if auto else "Manual"
    put(ks, r, 1, k["code"], "in", "0"); put(ks, r, 2, k["persp"], "in"); put(ks, r, 3, k["name"], "in", wrap=True)
    put(ks, r, 4, k["weight"], "in", PCT); put(ks, r, 5, mf, "in")
    for c, key, src in ((6, "fy", "fy_b"), (7, "fy_v2", "fy_v2"), (9, "ytd", "ytd_b"), (10, "ytd_v2", "ytd_v2"), (11, "act", "ytd_a")):
        if k["spend"]:
            put(ks, r, c, "=SUMIFS(%s,%s,$A%d)" % (pjr(src), pjr("kpi"), r), "f", KCOLS[c - 1][3])
        elif auto and c == 11:
            put(ks, r, c, auto[1].replace("{r}", str(r)), "wk", KCOLS[c - 1][3])
        elif auto and better == "Lower" and c == 9:
            put(ks, r, c, "=F%d" % r, "f", KCOLS[c - 1][3])
        else:
            put(ks, r, c, k[key] or 0, "in", KCOLS[c - 1][3])
    put(ks, r, 8, "=G%d-F%d" % (r, r), "f", KCOLS[7][3])
    put(ks, r, 12, "=K%d-I%d" % (r, r), "f", KCOLS[11][3])
    put(ks, r, 13, '=IFERROR(IF($R%d="Lower",IF(K%d<=I%d,1,I%d/K%d),MIN(K%d/I%d,1)),0)' % ((r,) * 7), "f", PCT)
    put(ks, r, 14, "=M%d*D%d" % (r, r), "f", '0.00%;(0.00%);"-"')
    put(ks, r, 15, '=IF(D%d=0,"Not weighted",IF(M%d>=0.95,"On target",IF(M%d>=0.8,"Watch","Below target")))' % (r, r, r), "f")
    put(ks, r, 16, k["remarks"], "in")
    put(ks, r, 17, auto[4] if auto else None, "in", "@")
    put(ks, r, 18, better, "in")
    how = auto[2] if auto else ("Monthly Plan: Σ YTD Actual of the projects with this KPI Code (Projects sheet)." if k["spend"] else
                                "Manual: not in the weekly report – type the actual (source: " + {"Safety": "SRGC", "Financial": "Finance"}.get(k["persp"].split()[0], "PMO / owner department") + ").")
    put(ks, r, 19, how, "f", wrap=True)
    ks.row_dimensions[r].height = 30
dv = DataValidation(type="list", formula1='"Higher,Lower"'); dv.add("R5:R200"); ks.add_data_validation(dv)
dv = DataValidation(type="list", formula1='"Monthly Plan,Weekly Report,Manual"'); dv.add("E5:E200"); ks.add_data_validation(dv)
add_table(ks, "KPI_Scorecard", "A4:S%d" % KN)
ks.cell(KN + 2, 3, "Total weight (must be 100%)").font = font(bold=True)
c = ks.cell(KN + 2, 4, "=SUM(D5:D%d)" % KN); c.number_format = PCT; c.font = font(bold=True)
ks.cell(KN + 2, 13, "Scorecard result").font = font(bold=True)
c = ks.cell(KN + 2, 14, "=SUM(N5:N%d)" % KN); c.number_format = '0.00%'; c.font = font(bold=True, color=TEAL, size=12)
KTOT = "'KPI Scorecard'!$N$%d" % (KN + 2)
ks.freeze_panes = "D5"

# ------------------------------------------------------------------ Delivery KPI
dk = S["Delivery KPI"]
title(dk, "Delivery KPI – target completion dates",
      "Type the project code and its target date. Everything else is read from Projects and the weekly files. Feeds KPI 12 (% of delivery against approved business plan): delivered ÷ due.")
DCOLS = [
    ("Project Code", "in", 10, "@", "Input"), ("Target Completion Date", "in", 13, DATE, "Input"), ("Remarks", "in", 30, None, "Input"),
    ("Project", "f", 40, None, "Calc: Projects"), ("Project Phase", "f", 12, None, "Calc: Projects"), ("Budget (Full Cost, SAR)", "f", 15, NUM, "Calc: Projects"),
    ("Forecast Finish (weekly)", "f", 13, DATE, "Calc: Projects ◂ weekly"), ("SPI (weekly)", "f", 8, "0.00", "Calc: Projects ◂ weekly"),
    ("Actual % (weekly)", "f", 9, PCT, "Calc: Projects ◂ weekly"), ("Days vs Target", "f", 9, '#,##0;(#,##0);"0"', "Calc: Forecast − Target"),
    ("Status", "f", 16, None, "Calc: from the dates"),
    ("Due by Report Date?", "f", 10, None, "Calc: Target ≤ report date"), ("Delivered?", "f", 10, None, "Calc: weekly Actual % = 100%"),
]
header_row(dk, 4, DCOLS)
def pjx(k, r): return 'IFERROR(INDEX(Projects!$%s$5:$%s$500,MATCH($A%d,Projects!$A$5:$A$500,0)),"")' % (PC[k], PC[k], r)
for i, x in enumerate(M["delivery"]):
    r = 5 + i
    put(dk, r, 1, x["code"], "in", "@"); put(dk, r, 2, ddate(x["target"]), "in", DATE); put(dk, r, 3, None, "in")
    put(dk, r, 4, "=" + pjx("name", r), "f"); put(dk, r, 5, "=" + pjx("phase", r), "f"); put(dk, r, 6, "=" + pjx("full_cost", r), "f", NUM)
    put(dk, r, 7, "=" + pjx("w_fin", r), "f", DATE); put(dk, r, 8, "=" + pjx("w_spi", r), "f", "0.00"); put(dk, r, 9, "=" + pjx("w_ac", r), "f", PCT)
    put(dk, r, 10, '=IF(ISNUMBER(G%d),G%d-B%d,"")' % (r, r, r), "f", '#,##0;(#,##0);"0"')
    put(dk, r, 11, '=IF(NOT(ISNUMBER(G%d)),"No weekly data",IF(G%d<=B%d,"On track","Late"))' % (r, r, r), "f")
    put(dk, r, 12, '=IF(B%d<=cfg_ReportDate,"Yes","No")' % r, "f")
    put(dk, r, 13, '=IF(N(I%d)>=0.999,"Yes","No")' % r, "f")
DN = 4 + len(M["delivery"])
add_table(dk, "Delivery_Targets", "A4:M%d" % DN)
dv = DataValidation(type="list", formula1="=Projects!$A$5:$A$500"); dv.add("A5:A200"); dk.add_data_validation(dv)

# ------------------------------------------------------------------ Data Checks (live)
dc = S["Data Checks"]
title(dc, "Data Checks – the plan (Projects) compared with the weekly project files",
      "Recalculates by itself. 'Mismatch' = the two sources disagree: fix the wrong one at its source (Projects row, or the project's weekly file and Refresh All).")
CCOLS = [
    ("Code", "f", 9, None, "Calc: Projects"), ("Project", "f", 16, None, "Calc: Projects"), ("In Weekly Report?", "f", 9, None, "Calc: weekly lookup"),
    ("Contract Value · Plan", "f", 14, NUM, "Projects"), ("Contract Value · Weekly", "wk", 14, NUM, "Contract Value"), ("CV Check", "f", 9, None, "± 1 SAR"),
    ("Start · Plan", "f", 12, DATE, "Projects"), ("Start · Weekly Baseline", "wk", 12, DATE, "Start Date \nBaseline"), ("Start Check", "f", 9, None, ""),
    ("End · Plan", "f", 12, DATE, "Projects"), ("End · Weekly Baseline", "wk", 12, DATE, "End Date\nBaseline"), ("End Check", "f", 9, None, ""),
    ("PO · Plan", "f", 9, "0", "Projects"), ("PO · Weekly", "wk", 9, "0", "(Con) PO Number"), ("PO Check", "f", 9, None, ""),
    ("Paid · Plan (ERP)", "f", 14, NUM, "Projects"), ("Paid · Weekly", "wk", 14, NUM, "Paid Amount"), ("Paid Check", "f", 9, None, "± 1 SAR"),
    ("Plan Sanity", "f", 34, None, "rules on the plan alone"), ("Result", "f", 22, None, "summary"),
]
header_row(dc, 4, CCOLS)
def pref(k, r): return "Projects!$%s%d" % (PC[k], r)
NCHK = 40
for i in range(NCHK):
    r, pr = 5 + i, 5 + i
    blank = 'IF($A%d="","",' % r
    put(dc, r, 1, '=IF(%s="","",%s)' % (pref("code", pr), pref("code", pr)), "f", "@")
    put(dc, r, 2, '=IF($A%d="","",%s)' % (r, pref("short", pr)), "f")
    put(dc, r, 3, '=IF($A%d="","",IF(%s="","No","Yes"))' % (r, pref("w_src", pr)), "f")
    for c, k in ((4, "cv"), (7, "start"), (10, "end"), (13, "po"), (16, "paid")):
        put(dc, r, c, '=IF(OR($A%d="",%s=""),"",%s)' % (r, pref(k, pr), pref(k, pr)), "f", CCOLS[c - 1][3])
    for c in (5, 8, 11, 14, 17):
        put(dc, r, c, '=IF($C%d<>"Yes","",%s' % (r, wk_lookup(pref("wcode", pr), "%s$3" % CL(c))[1:]) + ")", "wk", CCOLS[c - 1][3])
    for c, tol in ((6, 1), (9, 0), (12, 0), (18, 1)):
        a, b = CL(c - 2), CL(c - 1)
        put(dc, r, c, '=IF(OR($C{r}<>"Yes",{b}{r}=""),"",IF(ABS(N({a}{r})-N({b}{r}))<={t},"OK","Mismatch"))'.format(r=r, a=a, b=b, t=tol), "f")
    put(dc, r, 15, '=IF(OR($C{r}<>"Yes",N{r}=""),"",IF(TRIM(M{r}&"")=TRIM(N{r}&""),"OK","Mismatch"))'.format(r=r), "f")
    put(dc, r, 19, ('=IF($A{r}="","",_xlfn.TEXTJOIN("; ",TRUE,'
                    'IF(AND(N({cv})>N({fc}),N({fc})>0),"Contract value > full cost",""),'
                    'IF({st}="","No start date",""),'
                    'IF(AND(ISNUMBER({en}),{en}<cfg_ReportDate,OR({ph}="Execution",{ph}="On hold")),"End date passed but phase is "&{ph},""),'
                    'IF(AND({st}<>"",{en}<>"",N({en})<N({st})),"End before start","")))').format(
        r=r, cv=pref("cv", pr), fc=pref("full_cost", pr), st=pref("start", pr), en=pref("end", pr), ph=pref("phase", pr)), "f")
    put(dc, r, 20, '=IF($A{r}="","",IF(COUNTIF(D{r}:R{r},"Mismatch")>0,COUNTIF(D{r}:R{r},"Mismatch")&" mismatch(es) with weekly",'
                   'IF(S{r}<>"","Check plan data",IF($C{r}="Yes","OK","OK (no weekly file)"))))'.format(r=r), "f")
add_table(dc, "Data_Checks", "A4:T%d" % (4 + NCHK))
dc.freeze_panes = "C5"
red = PatternFill("solid", start_color="F8D7DA", end_color="F8D7DA"); green = PatternFill("solid", start_color="D4EDDA", end_color="D4EDDA")
dc.conditional_formatting.add("D5:T%d" % (4 + NCHK), FormulaRule(formula=['ISNUMBER(SEARCH("mismatch",D5))'], fill=red, font=Font(name=F, color="9C0006", bold=True)))
dc.conditional_formatting.add("S5:S%d" % (4 + NCHK), FormulaRule(formula=['S5<>""'], fill=fill("FFF3CD")))
dc.conditional_formatting.add("T5:T%d" % (4 + NCHK), FormulaRule(formula=['LEFT(T5,2)="OK"'], fill=green))

# static findings from the merge -> own sheet
rn = S["Review Notes"]
title(rn, "Review Notes – what was found while merging the two files (for discussion)",
      "Each line gives the exact place it came from (file › sheet › cells). This sheet can be deleted once the decisions are taken.")
FR = 4
FCOLS = [("#", 5), ("Topic", 24), ("Source (file › sheet › cells)", 44), ("What was found", 60), ("What the new file does", 48), ("Decision needed", 26)]
for j, (h, w) in enumerate(FCOLS):
    c = rn.cell(FR, 1 + j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(DARK); c.border = BOX
    c.alignment = Alignment(wrap_text=True, vertical="center"); rn.column_dimensions[CL(1 + j)].width = w
FIND = M.get("findings", [])
for i, row in enumerate(FIND):
    for j, v in enumerate([i + 1] + row):
        c = rn.cell(FR + 1 + i, 1 + j, v); c.font = font(size=9, bold=(j == 1)); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
        if j == 5 and v not in ("–", None): c.fill = fill("FFF3CD")
    rn.row_dimensions[FR + 1 + i].height = 52
add_table(rn, "Review_Notes", "A%d:F%d" % (FR, FR + len(FIND)))
rn.freeze_panes = "A5"

# ------------------------------------------------------------------ Data Sources
ds = S["Data Sources"]
title(ds, "Data Sources – where every table comes from and who updates it",
      "Click a sheet name to open it. 'Power Query' tables are rebuilt from the weekly project files by Data ▸ Refresh All: never type in them.")
SCOLS = [("#", 5), ("Sheet (click)", 26), ("Excel Table", 30), ("What it holds", 46), ("Type", 18), ("Comes from", 50), ("How to update", 38), ("Status", 22)]
for j, (h, w) in enumerate(SCOLS):
    c = ds.cell(4, 1 + j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(TEAL); c.border = BOX
    c.alignment = Alignment(wrap_text=True, vertical="center"); ds.column_dimensions[CL(1 + j)].width = w
FOLDER = r"D:\Final Dashboards\Weekly Reporty Updates\Weekly Report Updates"
PQ = "Power Query"
CAT = [
    ("Read Me", "–", "How the file works, colour code, update steps", "Guide", "–", "–", "Active"),
    ("Dashboard", "–", "One-page summary: tiles, portfolio table, monthly spend, scorecard", "Report (formulas)", "Projects, Monthly Plan, KPI Scorecard, Data Checks", "Nothing to type", "Active"),
    ("Projects", "Projects", "Master data: one row per project (codes, names, PO, cost, contract, WC, paid, dates, phase) + calculated plan and weekly figures", "Manual input + formulas", "Was: Delivery Plan › Spending Plan (G:Q repeated 12× per project), NSR Project Data, Abbreviation", "Type in the blue columns", "Active"),
    ("Monthly Plan", "Monthly_Plan", "Monthly budget plan, V0, V2, actual and contractor forecast per project", "Manual input + formulas", "Was: Delivery Plan › Spending Plan (S:AD)", "Type in the blue columns; add 12 rows for a new project", "Active"),
    ("KPI Scorecard", "KPI_Scorecard", "KPI weights, targets, actuals and scorecard result", "Manual input + formulas", "Was: Delivery Plan › KPI_Summary + KPI Projects_Data (rows 24-33)", "Type manual KPIs; spend KPIs calculate", "Active"),
    ("Delivery KPI", "Delivery_Targets", "Target completion dates vs weekly forecast finish", "Manual input + formulas", "Was: Delivery Plan › Delivery KPI", "Type code + target date", "Active"),
    ("Data Checks", "Data_Checks", "Live comparison plan vs weekly + plan sanity rules", "Formulas", "Projects + Weekly Report Updates", "Nothing to type – fix the source", "Active"),
    ("Review Notes", "Review_Notes", "Findings from the merge, with source cells and the decisions needed", "Notes", "Both old files", "Take the decisions, then delete the sheet", "Temporary"),
    ("Settings", "List_*, Glossary", "Report date, YTD cut-off, drop-down lists, glossary", "Manual input", "Was: Delivery Plan › List, Abbreviation (right table)", "Change the YTD cut-off monthly", "Active"),
    ("Weekly Report Updates", "Weekly_Report_Updates", "Project header data from each weekly file (dates, progress, money, HSE)", PQ, FOLDER + " › every project file › sheet Masterdata_Reporting", "Data ▸ Refresh All", "Active"),
    ("S-Curve", "S_Curve", "Weekly cumulative plan / actual / forecast %", PQ, FOLDER + " › every project file › S-Curve", "Data ▸ Refresh All", "Active"),
    ("Project_Milestones_Progress_Com", "Project_Milestones_Progress_Combine", "Milestone progress per project", PQ, FOLDER, "Data ▸ Refresh All", "Active"),
    ("Master_Plan_Combine", "Master_Plan_Combine", "Overall line per project (start, finish, progress)", PQ, FOLDER, "Data ▸ Refresh All", "Active"),
    ("Weekly Achievements", "Weekly_Acievements", "This week's achievements per project", PQ, FOLDER, "Data ▸ Refresh All", "Active"),
    ("Lookahead Activities ", "Lookahead_Activities", "Next 7 days activities", PQ, FOLDER, "Data ▸ Refresh All", "Active"),
    ("Deliverable Status", "Deliverable_Status", "Submittals planned / submitted / approved", PQ, FOLDER, "Data ▸ Refresh All", "Active"),
    ("Interim Payment Certificate", "Interim_Payment_Certificate", "IPC / VO amounts", PQ, FOLDER, "Data ▸ Refresh All", "Active"),
    ("Area of Concern", "Area_of_Concern", "Issues / concerns raised in the weekly files", PQ, FOLDER, "Data ▸ Refresh All", "Active"),
    ("PMO Issue", "Issue_register", "Issue register (PMO cards + weekly report issues)", "Manual input", "Typed in this sheet", "Type new issues at the end", "Active"),
    ("Sheet1", "Table7", "Map: PMO card name ↔ weekly file (Source.Name); used by Data Model + lookups", "Manual input", "Typed in this sheet", "Add a row for a new project file", "Active (proposed: merge into Projects)"),
    ("Sheet1 (2)", "MLS", "Same map as Sheet1 with the code column", "Manual input", "Typed in this sheet", "–", "Hidden – duplicate of Sheet1"),
    ("Project Milestones Progress ", "Project_Milestones_Progress", "Older version of the milestones query", PQ, FOLDER, "Data ▸ Refresh All", "Hidden – duplicate of _Com"),
    ("Master Plan", "Master_Plan", "Older overall-line table built with formulas", PQ + " + formulas", FOLDER, "Data ▸ Refresh All", "Hidden – duplicate of Master_Plan_Combine"),
    ("Old_PBI", "Old_PBI", "Old Power BI layout of the weekly header data", PQ, FOLDER, "Data ▸ Refresh All", "Hidden – legacy"),
]
for i, row in enumerate(CAT):
    r = 5 + i
    vals = [i + 1] + list(row)
    for j, v in enumerate(vals):
        c = ds.cell(r, 1 + j, v); c.font = font(size=9); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
    link(ds.cell(r, 2), row[0], "A1"); ds.cell(r, 2).font = font(size=9, color="0563C1", underline="single")
    if "Hidden" in row[6]:
        for j in range(8): ds.cell(r, 1 + j).fill = fill("F2F2F2")
    ds.row_dimensions[r].height = 36
CN = 4 + len(CAT)
add_table(ds, "Source_Catalogue", "A4:H%d" % CN)
r0 = CN + 3
ds.cell(r0 - 1, 1, "Weekly project files loaded now (from Weekly Report Updates – changes on Refresh All)").font = font(bold=True, size=12, color=TEAL)
for j, h in enumerate(("#", "Weekly file (Source.Name)", "Project Code", "Report Date", "Open the row")):
    c = ds.cell(r0, 1 + j, h); c.font = font(bold=True, color="FFFFFF"); c.fill = fill(DARK); c.border = BOX
for i in range(25):
    r = r0 + 1 + i
    put(ds, r, 1, i + 1, "f")
    put(ds, r, 2, '=IFERROR(IF(INDEX(%s!$A$1:$A$500,%d)=0,"",INDEX(%s!$A$1:$A$500,%d)),"")' % (WR, i + 2, WR, i + 2), "wk")
    put(ds, r, 3, '=IF(B%d="","",INDEX(%s,%d))' % (r, WCODE, i + 2), "wk", "@")
    put(ds, r, 4, '=IF(B%d="","",INDEX(%s!$B$1:$B$500,%d))' % (r, WR, i + 2), "wk", DATE)
    put(ds, r, 5, '=IF(B%d="","",HYPERLINK("#\'Weekly Report Updates\'!A%d","▸ row %d"))' % (r, i + 2, i + 2), "f")
ds.freeze_panes = "A5"

# ------------------------------------------------------------------ Dashboard
db = S["Dashboard"]
db.sheet_view.zoomScale = 90
db.column_dimensions["A"].width = 2
for c, w in zip("BCDEFGHIJKLMNO", (9, 20, 11, 14, 9, 10, 10, 10, 12, 14, 14, 10, 12, 26)):
    db.column_dimensions[c].width = w
db["B1"] = "EPBU · NSR Projects – Database & Dashboard"; db["B1"].font = font(size=18, bold=True, color=TEAL)
db["B2"] = '="Report date "&TEXT(cfg_ReportDate,"dd-mmm-yyyy")&"   ·   YTD to end of "&TEXT(cfg_YTDMonth,"mmm-yyyy")&"   ·   "&Settings!$B$9'
db["B2"].font = font(size=9, italic=True, color="595959")
db.row_dimensions[1].height = 28
TILES = [
    ("Projects", '=COUNTA(Projects!$A$5:$A$500)', "0", "Projects", '=COUNTIF(Projects!$%s$5:$%s$500,"Delayed")&" delayed (weekly)"' % (PC["w_st"], PC["w_st"])),
    ("Contract Value", "=SUM(Projects!$%s$5:$%s$500)" % (PC["cv"], PC["cv"]), MLN, "Projects", '="Paid "&TEXT(SUM(Projects!$%s$5:$%s$500)/1000000,"#,##0.0")&"M (ERP)"' % (PC["paid"], PC["paid"])),
    ("2026 Budget Plan", "=SUM(Projects!$%s$5:$%s$500)" % (PC["fy_b"], PC["fy_b"]), MLN, "Monthly Plan", '="V2 "&TEXT(SUM(Projects!$%s$5:$%s$500)/1000000,"#,##0.0")&"M"' % (PC["fy_v2"], PC["fy_v2"])),
    ("YTD Budget Plan", "=SUM(Projects!$%s$5:$%s$500)" % (PC["ytd_b"], PC["ytd_b"]), MLN, "Monthly Plan", '="to end of "&TEXT(cfg_YTDMonth,"mmm")'),
    ("YTD Actual", "=SUM(Projects!$%s$5:$%s$500)" % (PC["ytd_a"], PC["ytd_a"]), MLN, "Monthly Plan", '=TEXT(IFERROR(D5/B5*0+SUM(Projects!$%s$5:$%s$500)/SUM(Projects!$%s$5:$%s$500),0),"0%%")&" of YTD plan"' % (PC["ytd_a"], PC["ytd_a"], PC["ytd_b"], PC["ytd_b"])),
    ("Scorecard Result", "=" + KTOT, "0.0%", "KPI Scorecard", '="weights "&TEXT(\'KPI Scorecard\'!$D$%d,"0%%")' % (4 + len(M["kpis"]) + 2)),
    ("Data Checks", '=COUNTIF(\'Data Checks\'!$T$5:$T$200,"*mismatch*")+COUNTIF(\'Data Checks\'!$T$5:$T$200,"Check*")', "0", "Data Checks", '="projects to check"'),
]
for i, (lab, f, fmt, src, sub) in enumerate(TILES):
    c1 = 2 + i * 2
    db.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c1 + 1)
    db.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c1 + 1)
    db.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c1 + 1)
    db.merge_cells(start_row=7, start_column=c1, end_row=7, end_column=c1 + 1)
    a = db.cell(4, c1, lab); a.font = font(size=9, bold=True, color="FFFFFF"); a.fill = fill(TEAL); a.alignment = Alignment(horizontal="center")
    b = db.cell(5, c1, f); b.font = font(size=18, bold=True, color=DARK); b.number_format = fmt; b.alignment = Alignment(horizontal="center"); b.fill = fill(LIGHT)
    s = db.cell(6, c1, sub); s.font = font(size=8, color="595959"); s.alignment = Alignment(horizontal="center"); s.fill = fill(LIGHT)
    l = db.cell(7, c1, "Source: %s ▸" % src); link(l, src, "A1"); l.font = font(size=8, color="0563C1", underline="single"); l.alignment = Alignment(horizontal="center"); l.fill = fill(LIGHT)
    for rr in range(4, 8):
        db.cell(rr, c1 + 1).fill = fill(TEAL if rr == 4 else LIGHT)
db.row_dimensions[5].height = 30
# fix the YTD actual subtitle (simple formula)
db.cell(6, 10).value = '=TEXT(IFERROR(SUM(Projects!$%s$5:$%s$500)/SUM(Projects!$%s$5:$%s$500),0),"0%%")&" of YTD plan"' % (PC["ytd_a"], PC["ytd_a"], PC["ytd_b"], PC["ytd_b"])

# portfolio table
R0 = 9
db.cell(R0, 2, "Portfolio by project").font = font(bold=True, size=12, color=TEAL)
l = db.cell(R0, 6, "Source: Projects ▸ (click a code to open its row)"); link(l, "Projects", "A4"); l.font = font(size=8, color="0563C1", underline="single")
PORT = [("Code", None, None), ("Project", "short", None), ("Phase", "phase", None), ("Contract Value", "cv", NUM), ("Paid % (ERP)", None, PCT),
        ("Planned % (weekly)", "w_pl", PCT), ("Actual % (weekly)", "w_ac", PCT), ("Progress Var.", "w_var", PCT), ("Status (weekly)", "w_st", None),
        ("YTD Plan", "ytd_b", NUM), ("YTD Actual", "ytd_a", NUM), ("YTD Achieved", None, PCT), ("Forecast Finish", "w_fin", DATE), ("Data Checks", "chk", None)]
for j, (h, k, fmt) in enumerate(PORT):
    c = db.cell(R0 + 1, 2 + j, h); c.font = font(bold=True, color="FFFFFF", size=9); c.fill = fill(DARK)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
db.row_dimensions[R0 + 1].height = 30
NP = 30
for i in range(NP):
    r, pr = R0 + 2 + i, 5 + i
    code = "Projects!$A%d" % pr
    db.cell(r, 2, '=IF(%s="","",HYPERLINK("#\'Projects\'!A%d",%s))' % (code, pr, code))
    for j, (h, k, fmt) in enumerate(PORT[1:], start=3):
        if k: f = '=IF($B%d="","",Projects!$%s%d)' % (r, PC[k], pr)
        elif h.startswith("Paid"): f = '=IF($B%d="","",IFERROR(Projects!$%s%d/Projects!$%s%d,""))' % (r, PC["paid"], pr, PC["cv"], pr)
        else: f = '=IF(OR($B%d="",N(K%d)=0),"",L%d/K%d)' % (r, r, r, r)
        db.cell(r, j, f)
    for j, (h, k, fmt) in enumerate(PORT):
        c = db.cell(r, 2 + j); c.font = font(size=9, color="0563C1" if j == 0 else "000000", underline="single" if j == 0 else None)
        c.border = BOX
        if fmt: c.number_format = fmt
        if i % 2: c.fill = fill("F7FAFB")
RL = R0 + 1 + NP
db.conditional_formatting.add("J%d:J%d" % (R0 + 2, RL), CellIsRule(operator="equal", formula=['"Delayed"'], fill=red, font=Font(name=F, color="9C0006", bold=True)))
db.conditional_formatting.add("J%d:J%d" % (R0 + 2, RL), CellIsRule(operator="equal", formula=['"At Risk"'], fill=fill("FFF3CD"), font=Font(name=F, color="7F6000", bold=True)))
db.conditional_formatting.add("J%d:J%d" % (R0 + 2, RL), CellIsRule(operator="equal", formula=['"On Time"'], fill=green, font=Font(name=F, color="155724", bold=True)))
db.conditional_formatting.add("I%d:I%d" % (R0 + 2, RL), FormulaRule(formula=['AND(ISNUMBER(I%d),I%d<-0.05)' % (R0 + 2, R0 + 2)], font=Font(name=F, color="C00000", bold=True)))
db.conditional_formatting.add("O%d:O%d" % (R0 + 2, RL), FormulaRule(formula=['ISNUMBER(SEARCH("mismatch",O%d))' % (R0 + 2)], fill=red, font=Font(name=F, color="9C0006")))
db.conditional_formatting.add("O%d:O%d" % (R0 + 2, RL), FormulaRule(formula=['O%d="Check plan data"' % (R0 + 2)], fill=fill("FFF3CD")))
db.conditional_formatting.add("H%d:H%d" % (R0 + 2, RL), DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="63BE7B"))
db.freeze_panes = "A4"

# monthly spend + chart
R1 = RL + 3
db.cell(R1, 2, "Monthly spend 2026 – all projects").font = font(bold=True, size=12, color=TEAL)
l = db.cell(R1, 6, "Source: Monthly Plan ▸"); link(l, "Monthly Plan", "A4"); l.font = font(size=8, color="0563C1", underline="single")
MH = ["Month", "Budget Plan", "V2 Forecast", "Actual", "Contractor Fcst", "Budget (Cum)", "V2 (Cum)", "Actual (Cum)", "Fcst (Cum)"]
for j, h in enumerate(MH):
    c = db.cell(R1 + 1, 2 + j, h); c.font = font(bold=True, color="FFFFFF", size=9); c.fill = fill(DARK); c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = BOX
for i in range(12):
    r = R1 + 2 + i
    c = db.cell(r, 2, "=DATE(cfg_Year,%d,1)" % (i + 1)); c.number_format = "mmm"; c.font = font(size=9); c.border = BOX
    for j, src in enumerate("DFGH"):
        c = db.cell(r, 3 + j, "=SUMIFS(%s,%s,$B%d)" % (mp(src), mp("B"), r)); c.number_format = MLN; c.font = font(size=9); c.border = BOX
    for j in range(4):
        col = CL(3 + j)
        c = db.cell(r, 7 + j, "=SUM(%s$%d:%s%d)" % (col, R1 + 2, col, r)); c.number_format = MLN; c.font = font(size=9); c.border = BOX
ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.title = "Monthly spend 2026 (SAR)"; ch.y_axis.title = None
ch.y_axis.numFmt = '#,##0,,"M"'; ch.height = 7.5; ch.width = 17
ch.add_data(Reference(db, min_col=3, max_col=6, min_row=R1 + 1, max_row=R1 + 13), titles_from_data=True)
ch.set_categories(Reference(db, min_col=2, min_row=R1 + 2, max_row=R1 + 13))
for s, col in zip(ch.series, ("BFBFBF", "7FBFCA", TEAL, "F4B183")):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ch.legend.position = "b"
db.add_chart(ch, "K%d" % (R1 + 1))

# scorecard mini table + progress chart
R2 = R1 + 17
db.cell(R2, 2, "KPI Scorecard").font = font(bold=True, size=12, color=TEAL)
l = db.cell(R2, 6, "Source: KPI Scorecard ▸"); link(l, "KPI Scorecard", "A4"); l.font = font(size=8, color="0563C1", underline="single")
for j, h in enumerate(("KPI", "Objective / KPI", "", "", "", "Weight", "% Achieved", "Result", "Indicator")):
    c = db.cell(R2 + 1, 2 + j, h); c.font = font(bold=True, color="FFFFFF", size=9); c.fill = fill(DARK); c.border = BOX
db.merge_cells(start_row=R2 + 1, start_column=3, end_row=R2 + 1, end_column=6)
for i in range(len(M["kpis"])):
    r, kr = R2 + 2 + i, 5 + i
    db.cell(r, 2, "='KPI Scorecard'!A%d" % kr)
    db.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    db.cell(r, 3, "='KPI Scorecard'!C%d" % kr)
    db.cell(r, 7, "='KPI Scorecard'!D%d" % kr).number_format = PCT
    db.cell(r, 8, "='KPI Scorecard'!M%d" % kr).number_format = PCT
    db.cell(r, 9, "='KPI Scorecard'!N%d" % kr).number_format = '0.00%;(0.00%);"-"'
    db.cell(r, 10, "='KPI Scorecard'!O%d" % kr)
    for cc in range(2, 11):
        db.cell(r, cc).font = font(size=9); db.cell(r, cc).border = BOX
RT = R2 + 2 + len(M["kpis"])
db.cell(RT, 8, "Total").font = font(bold=True)
c = db.cell(RT, 9, "=" + KTOT); c.number_format = "0.00%"; c.font = font(bold=True, color=TEAL)
db.conditional_formatting.add("J%d:J%d" % (R2 + 2, RT - 1), CellIsRule(operator="equal", formula=['"Below target"'], font=Font(name=F, color="C00000", bold=True)))
db.conditional_formatting.add("J%d:J%d" % (R2 + 2, RT - 1), CellIsRule(operator="equal", formula=['"On target"'], font=Font(name=F, color="2E7D32", bold=True)))
ch2 = BarChart(); ch2.type = "bar"; ch2.grouping = "clustered"; ch2.title = "Weekly progress – planned vs actual"
ch2.height = 9; ch2.width = 17; ch2.x_axis.scaling.orientation = "maxMin"; ch2.y_axis.numFmt = "0%"; ch2.y_axis.scaling.max = 1; ch2.y_axis.scaling.min = 0
NPJ = len(M["projects"])
ch2.add_data(Reference(db, min_col=7, max_col=8, min_row=R0 + 1, max_row=R0 + 1 + NPJ), titles_from_data=True)
ch2.set_categories(Reference(db, min_col=3, min_row=R0 + 2, max_row=R0 + 1 + NPJ))
for s, col in zip(ch2.series, ("BFBFBF", TEAL)):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ch2.legend.position = "b"
db.add_chart(ch2, "K%d" % (R2 + 1))

# ------------------------------------------------------------------ Read Me
rm = S["Read Me"]
rm.column_dimensions["A"].width = 3; rm.column_dimensions["B"].width = 30; rm.column_dimensions["C"].width = 100
rm["B1"] = "EPBU · NSR Projects Database"; rm["B1"].font = font(size=20, bold=True, color=TEAL)
rm["B2"] = "One workbook for the weekly progress (Power Query) and the 2026 delivery plan, KPIs and spending (manual input)."; rm["B2"].font = font(size=10, color="595959")
rm["B3"] = "=Settings!$B$9"; rm["B3"].font = font(size=10, bold=True, color="C00000")
blocks = [
    ("How the file works", [
        ("1 · Weekly data", "The green-tab sheets (Weekly Report Updates, S-Curve, Milestones, …) are built by Power Query from every project file in the weekly folder. Do not type in them. Data ▸ Refresh All updates them."),
        ("2 · Plan data", "The teal-tab sheets are typed by the cost control team: Projects (one row per project), Monthly Plan (one row per project per month), KPI Scorecard, Delivery KPI, Settings."),
        ("3 · Reports", "Dashboard, Data Checks and Data Sources only calculate. They link the plan and the weekly data through the project code (Projects ▸ Weekly Report Code = weekly file 'Project Code')."),
        ("4 · Find the source", "Every report cell has a source: click a tile's 'Source ▸' link, a project code on the Dashboard, or a sheet name in Data Sources. Row 3 of every input sheet says where each column comes from."),
    ]),
    ("Colour code", [
        ("Blue text, yellow cell", "Input: type here."),
        ("Bright yellow cell", "Key setting (Settings ▸ YTD cut-off month)."),
        ("Black text", "Formula: do not type over it."),
        ("Green text", "Read from the weekly project files (Power Query)."),
        ("Teal / dark / green header", "Input column / calculated column / weekly column."),
    ]),
    ("Weekly update (every week)", [
        ("Step 1", "Put the new weekly project files in the weekly folder (same file names as before)."),
        ("Step 2", "Data ▸ Refresh All. Wait until the status bar is clear."),
        ("Step 3", "Open Data Checks: every project should read OK. A 'Mismatch' means the plan (Projects) and the weekly file disagree: fix the wrong one."),
        ("Step 4", "Save. Upload this file to the website (Import data) or link it to Power BI."),
    ]),
    ("Monthly update (after month close)", [
        ("Step 1", "Monthly Plan: type the month's Actual Spend (and update V2 / contractor forecast if they changed)."),
        ("Step 2", "Settings: move the YTD cut-off month to the closed month."),
        ("Step 3", "KPI Scorecard: type the manual KPI actuals (safety, governance, local content, …)."),
    ]),
    ("Adding a project", [
        ("Projects", "Add a row at the end of the table (the table grows by itself). Code must be unique. If the weekly file uses another code, put it in 'Weekly Report Code'."),
        ("Monthly Plan", "Add 12 rows (one per month) with the same code."),
        ("Weekly", "Put the project's weekly file in the weekly folder and Refresh All. Add it to Sheet1 (PMO card ↔ file name)."),
    ]),
    ("Power BI / website", [
        ("Tables", "Every sheet is an Excel table with one header row: Projects, Monthly_Plan, KPI_Scorecard, Delivery_Targets, Data_Checks + the weekly query tables. Join them on Code / Project Code."),
        ("Website", "The NSR dashboard site reads this file as both the 'PBI Weekly Report' and the 'Delivery Plan'. The old two files still work."),
    ]),
]
r = 5
for head, rows in blocks:
    c = rm.cell(r, 2, head); c.font = font(bold=True, size=12, color="FFFFFF"); c.fill = fill(TEAL)
    rm.cell(r, 3).fill = fill(TEAL); r += 1
    for k, v in rows:
        a = rm.cell(r, 2, k); a.font = font(bold=True, size=10); a.alignment = Alignment(vertical="top"); a.border = BOX
        b = rm.cell(r, 3, v); b.font = font(size=10); b.alignment = Alignment(wrap_text=True, vertical="top"); b.border = BOX
        if head == "Colour code":
            if k.startswith("Blue"): b.font = FT_IN; b.fill = FILL_IN
            elif k.startswith("Bright"): b.fill = FILL_KEY
            elif k.startswith("Green"): b.font = FT_LINK
        rm.row_dimensions[r].height = 30
        r += 1
    r += 1
c = rm.cell(r, 2, "Sheets"); c.font = font(bold=True, size=12, color="FFFFFF"); c.fill = fill(TEAL); rm.cell(r, 3).fill = fill(TEAL); r += 1
for n in ORDER[1:]:
    a = rm.cell(r, 2, n); link(a, n, "A1"); a.border = BOX
    rm.cell(r, 3, {"Dashboard": "One-page summary", "Projects": "Project master data (input)", "Monthly Plan": "Monthly spend plan and actuals (input)",
                   "KPI Scorecard": "KPIs and scorecard result (input + calc)", "Delivery KPI": "Target completion dates (input)",
                   "Data Checks": "Plan vs weekly comparison (live)", "Review Notes": "Merge findings and decisions needed (temporary)", "Data Sources": "Where every table comes from",
                   "Settings": "Report date, YTD cut-off, lists, glossary"}[n]).border = BOX
    rm.cell(r, 3).font = font(size=10); r += 1
a = rm.cell(r, 2, "Weekly query sheets"); link(a, "Data Sources", "A4"); a.border = BOX
rm.cell(r, 3, "See Data Sources for the full list").border = BOX

for n in ORDER:
    w = S[n]
    w.page_setup.orientation = "landscape"; w.page_setup.paperSize = w.PAPERSIZE_A4
    w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0; w.sheet_properties.pageSetUpPr.fitToPage = True
    w.print_options.horizontalCentered = True
    w.page_margins.left = w.page_margins.right = 0.4; w.page_margins.top = w.page_margins.bottom = 0.5

# ------------------------------------------------------------------ weekly copy for recalc
src = openpyxl.load_workbook(wk_p, data_only=True)["Weekly Report Updates"]
cp = wb.create_sheet("Weekly Report Updates")
for row in src.iter_rows():
    for c in row:
        if c.value is not None: cp.cell(c.row, c.column, c.value)

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != F and c.value is not None:
                c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic, color=c.font.color, underline=c.font.underline)
wb.save(out_p)
print("saved", out_p, {n: S[n].max_row for n in ORDER}, "projects col map:", PC["chk"], PC["w_src"])
json.dump({"PC": PC, "KTOT": KTOT, "R0": R0}, open(out_p + ".map.json", "w"))
