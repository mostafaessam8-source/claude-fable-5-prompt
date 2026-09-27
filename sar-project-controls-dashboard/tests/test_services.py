import time

import pytest

from conftest import load_text
from sar_pcd.analysis.filters import FilterState
from sar_pcd.analysis.health import FAIL, NOT_PERFORMED, PASS
from sar_pcd.analysis.risk import RiskImportError, heatmap, load_register
from sar_pcd.analysis.settings import AnalysisSettings
from sar_pcd.mapping.profile import default_profile
from sar_pcd.services.dashboard import ContextCache, build
from sar_pcd.services.pipeline import build_context
from sar_pcd.services.workspace import Workspace


def test_filters_recompute(demo_ctx):
    full = build(demo_ctx)
    proc = build(demo_ctx, FilterState(phase="Procurement"))
    assert 0 < len(proc.rows) < len(full.rows)
    assert all(r.phase == "Procurement" for r in proc.rows)
    assert proc.kpis["progress_actual"].value != full.kpis["progress_actual"].value
    loc = build(demo_ctx, FilterState(location="Section B (km 40-85)"))
    assert all(r.dims.get("Location") == "Section B (km 40-85)" for r in loc.rows)
    s = build(demo_ctx, FilterState(search="track installation"))
    assert {r.name for r in s.rows} == {"Track Installation - Section A", "Track Installation - Section B",
                                        "Track Installation - Section C"}


def test_kpis_are_auditable(demo_ctx):
    d = build(demo_ctx)
    for key in ("progress_actual", "progress_planned", "spi", "cpi", "cp_critical"):
        k = d.kpis[key]
        assert k.formula and (k.source or k.notes)
    assert len(d.kpis["cp_critical"].keys) == d.kpis["cp_critical"].value


def test_health_checks_show_threshold_and_population(demo_ctx):
    checks = ContextCache(demo_ctx).health
    ids = {c.id for c in checks}
    assert {"DCMA-01", "DCMA-06", "DCMA-07", "DCMA-14", "Q-03", "Q-07"} <= ids
    for c in checks:
        assert c.result in (PASS, FAIL, "Info", NOT_PERFORMED)
        if c.result in (PASS, FAIL):
            assert c.threshold
    neg = next(c for c in checks if c.id == "DCMA-07")
    assert neg.count == len(neg.keys) > 0 and neg.result == FAIL
    assert next(c for c in checks if c.id == "DCMA-12").result == NOT_PERFORMED


def test_validation_findings_listed(demo_ctx):
    findings, score = ContextCache(demo_ctx).validation
    cats = {f.category for f in findings}
    assert "Mapping Confidence" in cats and "Missing Baseline Activities" in cats
    assert 0 <= score <= 100


def test_risk_register_import_and_no_fabrication(tmp_path, demo_ctx):
    d = build(demo_ctx)
    assert d.risk["available"] is False and d.risk["message"] == "Risk Data Not Available"
    p = tmp_path / "risks.csv"
    p.write_text("Risk ID,Title,Probability,Impact,Status\nR1,Late rails,4,5,Open\nR2,Ground,2,2,Open\n"
                 "R3,Closed one,5,5,Closed\nR4,Bad,9,1,Open\n")
    items, warnings = load_register(p)
    assert len(items) == 3 and len(warnings) == 1
    hm = heatmap(items, AnalysisSettings().risk_bands)
    assert hm["total"] == 2 and hm["grid"][3][4] == 1 and hm["counts"]["Very High"] == 1
    bad = tmp_path / "bad.csv"
    bad.write_text("a,b\n1,2\n")
    with pytest.raises(RiskImportError):
        load_register(bad)


def test_workspace_roundtrip_and_audit(tmp_path, variants):
    bl, cu = variants["B"]
    b, c = load_text(bl), load_text(cu)
    ws = Workspace.create(tmp_path / "p.sarpcd", "Proj B")
    ws.save_schedule("baseline", b, bl.encode(), "bl.xer")
    ws.save_schedule("current", c, cu.encode(), "cu.xer")
    prof = ws.load_profile()
    prof.activity_overrides["D-1010"] = "Handover"
    ws.save_profile(prof, note="override D-1010")
    s = ws.load_settings()
    s.near_critical_days = 15
    ws.save_settings(s)
    ws.close()
    ws = Workspace(tmp_path / "p.sarpcd")
    ctx = build_context(ws.load_schedule("baseline"), ws.load_schedule("current"), ws.load_profile(), ws.load_settings())
    assert ctx.settings.near_critical_days == 15
    r = next(r for r in ctx.rows if r.code == "D-1010")
    assert r.phase == "Handover" and r.phase_level == "Override"
    actions = [a["action"] for a in ws.audit_trail()]
    assert "Save Mapping Profile" in actions and "Change Settings" in actions and "Create Project" in actions
    dup = ws.duplicate(tmp_path / "copy.sarpcd", "Copy")
    assert dup.name == "Copy" and dup.load_schedule("current") is not None
    dup.archive()
    assert dup.archived


def test_performance_10k_activities():
    """10k activities: parse + normalize + match + classify + KPIs in reasonable time."""
    from datetime import date
    from sar_pcd.demo.generator import GA, GenProject, build_xer
    n = 10_000
    wbs = [(f"W{i}", None, str(i), name) for i, name in enumerate(
        ["Engineering", "Procurement", "Construction", "Testing & Commissioning", "Handover"])]
    names = ["Design package", "Purchase order", "Concrete works", "System testing", "As-built docs"]
    acts = []
    for i in range(n):
        g = i * 5 // n
        preds = [(f"A{i - 1}", "FS", 0)] if i % 50 else []
        if i >= 50 and i % 50 == 0:
            preds = [(f"A{i - 50}", "SS", 2)]
        acts.append(GA(f"A{i}", f"{names[g]} {i}", f"W{g}", 3 + i % 7, preds, rate=100.0))
    proj = GenProject("BIG", "Big", date(2024, 1, 7), wbs, acts, {}, "monfri")
    t0 = time.time()
    bl = build_xer(proj, proj_id=1, task_base=1)
    cu = build_xer(proj, proj_id=2, task_base=100000)
    gen = time.time() - t0
    t0 = time.time()
    ctx = build_context(load_text(bl), load_text(cu), default_profile(), AnalysisSettings())
    d = build(ctx)
    elapsed = time.time() - t0
    assert len(ctx.rows) == n
    assert d.kpis["progress_planned"].available
    assert elapsed < 60, f"analysis took {elapsed:.1f}s (generation {gen:.1f}s)"


def test_dynamic_excel_has_live_formulas(tmp_path, demo_ctx):
    from openpyxl import load_workbook
    from sar_pcd.reporting.excel_dynamic import COL, export_dynamic
    d = build(demo_ctx)
    path = tmp_path / "dyn.xlsx"
    info = export_dynamic(path, d)
    assert info["rows"] == len(demo_ctx.rows) and info["formulas"] > 5 * info["rows"]
    wb = load_workbook(path)
    assert {"Dashboard", "Data", "S-Curve", "Calendars", "Lists", "Read Me"} <= set(wb.sheetnames)
    for name in ("DataDate", "F_Phase", "F_Disc", "D_Inc", "D_PF", "K_Actual", "K_Planned", "CalCum", "CalOrigin"):
        assert name in wb.defined_names
    dash = wb["Dashboard"]
    assert dash["A16"].value.startswith("=IFERROR(SUMPRODUCT(D_Inc,D_InCur,D_W,D_Pct)")
    assert "EV/K_AC" in dash["B19"].value.replace(" ", "") or "K_EV/K_AC" in dash["B19"].value
    data = wb["Data"]
    assert data[f"{COL['pf']}2"].value.startswith("=IF(")
    assert "INDEX(CalCum" in data[f"{COL['pf']}2"].value and "_xlfn" not in data[f"{COL['pf']}2"].value
    assert len(dash._charts) >= 3
