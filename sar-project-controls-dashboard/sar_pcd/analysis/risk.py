"""Risk register support.

Risks are NEVER derived from schedule activities. A 5x5 heatmap is produced
only from an explicit risk register (CSV or Excel) imported by the user.
Risk tables inside an XER (e.g. PROJRISK) are detected and reported, but their
field meaning is version-dependent and not validated, so they are not used.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from ..mapping.text import normalize

ALIASES = {
    "id": ["risk id", "id", "ref", "risk ref", "risk no", "risk number"],
    "title": ["title", "risk", "risk title", "description", "risk description", "name"],
    "probability": ["probability", "likelihood", "p", "prob"],
    "impact": ["impact", "consequence", "severity", "i"],
    "owner": ["owner", "risk owner", "responsible"],
    "status": ["status", "state"],
    "category": ["category", "type", "discipline"],
    "activity": ["activity id", "linked activity", "activity"],
    "response": ["response", "mitigation", "treatment", "action"],
}


@dataclass
class RiskItem:
    id: str
    title: str
    probability: int
    impact: int
    owner: str = ""
    status: str = "Open"
    category: str = ""
    activity: str = ""
    response: str = ""

    @property
    def score(self) -> int:
        return self.probability * self.impact


def level(score: int, bands: dict) -> str:
    for name in ("Low", "Medium", "High", "Very High"):
        if score <= bands.get(name, 25):
            return name
    return "Very High"


class RiskImportError(Exception):
    pass


def _rows_from_file(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "") for h in next(it)]
        return [dict(zip(header, r)) for r in it if any(v not in (None, "") for v in r)]
    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(text.splitlines()))


def load_register(path: str | Path) -> tuple[list[RiskItem], list[str]]:
    path = Path(path)
    rows = _rows_from_file(path)
    if not rows:
        raise RiskImportError("The risk register is empty.")
    cols = {normalize(c): c for c in rows[0].keys()}
    mapping = {}
    for field_name, names in ALIASES.items():
        for n in names:
            if normalize(n) in cols:
                mapping[field_name] = cols[normalize(n)]
                break
    for req in ("probability", "impact"):
        if req not in mapping:
            raise RiskImportError(f"Column for '{req}' not found. Expected one of: {', '.join(ALIASES[req])}.")
    items, warnings = [], []
    for i, r in enumerate(rows, start=2):
        try:
            p = int(float(r[mapping["probability"]]))
            im = int(float(r[mapping["impact"]]))
        except (TypeError, ValueError):
            warnings.append(f"Row {i}: probability/impact not numeric - skipped.")
            continue
        if not (1 <= p <= 5 and 1 <= im <= 5):
            warnings.append(f"Row {i}: probability/impact outside 1-5 - skipped.")
            continue
        g = lambda k: str(r.get(mapping.get(k, ""), "") or "")
        items.append(RiskItem(g("id") or f"R{i - 1:03d}", g("title"), p, im, g("owner"), g("status") or "Open",
                              g("category"), g("activity"), g("response")))
    return items, warnings


def heatmap(items: list[RiskItem], bands: dict, only_open: bool = True) -> dict:
    grid = [[0] * 5 for _ in range(5)]   # grid[prob-1][impact-1]
    counts = {"Low": 0, "Medium": 0, "High": 0, "Very High": 0}
    used = [x for x in items if not only_open or x.status.lower() not in ("closed", "retired")]
    for x in used:
        grid[x.probability - 1][x.impact - 1] += 1
        counts[level(x.score, bands)] += 1
    return {"grid": grid, "counts": counts, "total": len(used)}
